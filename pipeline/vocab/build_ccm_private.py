# -*- coding: utf-8 -*-
"""Build a PRIVATE SKOS file of CSA Cloud Controls Matrix v4.1 controls from a licensed copy (decision D2, 2026-10-09).

The CCM terms do not allow redistribution, so the input and the output stay out of Git:
    input : data/private/ (any of: CSA OSCAL catalogue JSON, CCM JSON/YAML bundle, or the CCM .xlsx)
    output: ontology/vocab-private/ccm41-controls.ttl
Both folders are in .gitignore. The public file ontology/vocab/ccm41.ttl keeps the 17 domains; controls in the
private file use the same namespace and point to those domains with skos:broader, so mappings work when both are loaded.

Usage: python3 pipeline/vocab/build_ccm_private.py [path-to-bundle-file]
"""
import glob, json, os, re, sys
from rdflib import Graph, Namespace, Literal, URIRef
from rdflib.namespace import RDF, SKOS, DCTERMS, OWL, XSD

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ids = dict(l.rstrip("\n").split("\t") for l in open(os.path.join(ROOT, "src", "ids.tsv"), encoding="utf-8"))
C = lambda k: URIRef("https://w3id.org/crgo/CRGO_%07d" % int(ids[k]))
NS = Namespace("https://w3id.org/crgo/vocab/ccm41/")
ID_RE = re.compile(r"^([A-Z&]{2,4})-(\d{2})$")
safe = lambda i: re.sub(r"[^A-Za-z0-9._-]", "", i.replace("&", "and"))

def from_oscal(d):
    out = []
    def walk(controls):
        for c in controls:
            label = next((p["value"] for p in c.get("props", []) if p.get("name") == "label"), c.get("id", ""))
            st = next((p.get("prose") for p in c.get("parts", []) if p.get("name") == "statement"), None)
            out.append((label.upper(), c.get("title", ""), st))
            walk(c.get("controls", []))
    for g in d["catalog"].get("groups", []):
        walk(g.get("controls", []))
    walk(d["catalog"].get("controls", []))
    return out

def from_json_any(d):
    """Generic JSON/YAML bundle: find dicts that carry a CCM-style id and a title."""
    out = []
    def walk(x):
        if isinstance(x, dict):
            keys = {k.lower().replace(" ", "_"): k for k in x}
            idk = next((keys[k] for k in ("control_id", "id", "controlid", "ref") if k in keys), None)
            tk = next((keys[k] for k in ("control_title", "title", "name") if k in keys), None)
            if idk and tk and isinstance(x[idk], str) and ID_RE.match(x[idk].strip().upper()):
                sk = next((keys[k] for k in ("control_specification", "specification", "description", "statement") if k in keys), None)
                out.append((x[idk].strip().upper(), str(x[tk]).strip(), str(x[sk]).strip() if sk else None))
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
    walk(d)
    return out

def from_xlsx(path):
    from openpyxl import load_workbook
    out = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        hdr = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if hdr is None:
                low = [c.lower() for c in cells]
                if "control id" in low:
                    hdr = {"id": low.index("control id"),
                           "title": next((i for i, c in enumerate(low) if c == "control title"), None),
                           "spec": next((i for i, c in enumerate(low) if c.startswith("control specification")), None)}
                continue
            cid = cells[hdr["id"]].upper() if hdr["id"] < len(cells) else ""
            if ID_RE.match(cid):
                title = cells[hdr["title"]] if hdr["title"] is not None else ""
                spec = cells[hdr["spec"]] if hdr["spec"] is not None else None
                out.append((cid, title, spec or None))
        if out:
            break
    return out

def load(path):
    if path.lower().endswith(".xlsx"):
        return from_xlsx(path), "xlsx"
    if path.lower().endswith((".yaml", ".yml")):
        import yaml
        d = yaml.safe_load(open(path, encoding="utf-8"))
    else:
        d = json.load(open(path, encoding="utf-8"))
    if isinstance(d, dict) and "catalog" in d:
        return from_oscal(d), "OSCAL"
    return from_json_any(d), "JSON/YAML"

def main():
    paths = sys.argv[1:] or sorted(glob.glob(os.path.join(ROOT, "data", "private", "*")))
    paths = [p for p in paths if p.lower().endswith((".json", ".yaml", ".yml", ".xlsx"))]
    if not paths:
        sys.exit("No CCM bundle found. Put the licensed CCM v4.1 file (OSCAL JSON, JSON/YAML or .xlsx) in data/private/.")
    rows, kind, src = [], None, None
    for p in paths:
        rows, kind = load(p)
        if rows:
            src = p; break
    if not rows:
        sys.exit("No CCM control rows recognised in: " + ", ".join(paths))
    seen = {}
    for cid, title, spec in rows:
        seen.setdefault(cid, (title, spec))
    domains = {}
    pub = Graph(); pub.parse(os.path.join(ROOT, "ontology", "vocab", "ccm41.ttl"))
    for c in pub.subjects(RDF.type, SKOS.Concept):
        domains[str(pub.value(c, SKOS.notation))] = c
    g = Graph()
    for p, n in [("skos", SKOS), ("dcterms", DCTERMS), ("crgo", "https://w3id.org/crgo/"), ("ccm41", NS), ("owl", OWL)]:
        g.bind(p, n)
    o = URIRef("https://w3id.org/crgo/vocab/ccm41-controls")
    g.add((o, RDF.type, OWL.Ontology))
    g.add((o, DCTERMS.title, Literal("CSA Cloud Controls Matrix v4.1 controls (private, licensed copy)", lang="en")))
    g.add((o, DCTERMS.rights, Literal("(c) Cloud Security Alliance. Built from the project owner's licensed copy for internal, "
                                      "non-commercial use. Do not publish or redistribute.", lang="en")))
    g.add((o, DCTERMS.source, Literal("%s (%s)" % (os.path.basename(src), kind))))
    g.add((o, DCTERMS.date, Literal("2026-10-09", datatype=XSD.date)))
    missing_domain = set()
    for cid, (title, spec) in sorted(seen.items()):
        c = NS[safe(cid)]
        dom = ID_RE.match(cid).group(1)
        g.add((c, RDF.type, SKOS.Concept)); g.add((c, SKOS.inScheme, URIRef("https://w3id.org/crgo/vocab/ccm41")))
        g.add((c, SKOS.notation, Literal(cid))); g.add((c, SKOS.prefLabel, Literal(title or cid, lang="en")))
        if spec:
            g.add((c, SKOS.definition, Literal(spec, lang="en")))
        if dom in domains:
            g.add((c, SKOS.broader, domains[dom]))
        else:
            missing_domain.add(dom)
        g.add((c, C("governanceStatus"), C("StatusProposed"))); g.add((c, C("confidence"), C("ConfHigh")))
        g.add((c, C("evidenceSource"), Literal("CSA CCM v4.1, licensed copy: " + os.path.basename(src), lang="en")))
        g.add((c, C("proposedBy"), C("AgentClaude"))); g.add((c, C("introducedIn"), Literal("0.2.0")))
    out = os.path.join(ROOT, "ontology", "vocab-private"); os.makedirs(out, exist_ok=True)
    g.serialize(os.path.join(out, "ccm41-controls.ttl"), format="turtle")
    by_dom = {}
    for cid in seen:
        by_dom[ID_RE.match(cid).group(1)] = by_dom.get(ID_RE.match(cid).group(1), 0) + 1
    print("%d controls from %s (%s) -> ontology/vocab-private/ccm41-controls.ttl" % (len(seen), os.path.basename(src), kind))
    print("per domain: " + ", ".join("%s %d" % kv for kv in sorted(by_dom.items())))
    if missing_domain:
        print("WARNING: domains not in the public file: " + ", ".join(sorted(missing_domain)))

if __name__ == "__main__":
    main()
