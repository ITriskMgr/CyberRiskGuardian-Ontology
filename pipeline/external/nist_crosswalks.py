# -*- coding: utf-8 -*-
"""0.8.0: NIST crosswalks supplied by the project owner (downloaded from csrc.nist.gov into data/raw/external/nist/).

1. CSF 2.0 subcategory -> SP 800-53 Rev. 5 control, DERIVED: NIST's mapping of CSF 1.1 to SP 800-53r5
   (csf-pf-to-sp800-53r5-mappings.xlsx, sheet "CSF to SP 800-53r5") carried to CSF 2.0 through the CSF 2.0 export of
   NIST's Cybersecurity and Privacy Reference Tool (csf-export.json): a 1.1 subcategory kept in 2.0 keeps its controls; a
   withdrawn one passes them to the subcategories it was "incorporated into" or "moved to". Control enhancements are
   reduced to their base control (CRGO's SP 800-53 vocabulary holds base controls).
2. SP 800-53 Rev. 5 control -> ISO/IEC 27002:2022 control, from NIST's OLIR mapping of SP 800-53r5 to ISO/IEC 27001:2022
   (sp800-53r5-to-iso-27001-mapping-2022-OLIR-2023-10-12-UPDATED.xlsx): only references to Annex A controls (A.5 to A.8),
   which are the ISO/IEC 27002:2022 controls; references to the management-system clauses 4 to 10 are left out.
Writes data/library/nist_csf2_80053.csv and data/library/nist_80053_iso27002.csv."""
import csv, json, os, re, collections
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
N = os.path.join(ROOT, "data", "raw", "external", "nist")
base = lambda c: re.sub(r"-0*(\d+)", r"-\1", re.sub(r"\(.*\)", "", c.strip()))
def two(sub):  # ID.AM-1 -> ID.AM-01
    m = re.match(r"([A-Z]{2}\.[A-Z]{2})-(\d+)", sub)
    return "%s-%02d" % (m.group(1), int(m.group(2))) if m else None
d = json.load(open(os.path.join(N, "csf-export.json")))["response"]["elements"]
csf2 = {e["element_identifier"] for e in d["elements"] if e["element_type"] == "subcategory"}
moved = {}
for e in d["elements"]:
    if e["element_type"] == "withdraw_reason":
        moved[e["element_identifier"][3:]] = re.findall(r"[A-Z]{2}\.[A-Z]{2}-\d{2}", e["text"])
wb = openpyxl.load_workbook(os.path.join(N, "csf-pf-to-sp800-53r5-mappings.xlsx"), read_only=True)
m11 = collections.defaultdict(set)
for r in wb["CSF to SP 800-53r5"].iter_rows(values_only=True):
    sub, ctl = r[2], r[3]
    if isinstance(sub, str) and two(sub.split(":")[0]) and isinstance(ctl, str):
        for c in re.split(r"[,;]\s*", ctl):
            if re.match(r"[A-Z]{2}-\d+", c.strip()): m11[two(sub.split(":")[0])].add(base(c))
out1, stats = set(), collections.Counter()
for s11, ctls in m11.items():
    targets = moved.get(s11) if s11 in moved else ([s11] if s11 in csf2 else [])
    stats["kept" if s11 in csf2 and s11 not in moved else ("moved" if targets else "lost")] += 1
    for t in targets or []:
        for c in ctls: out1.add((t, c, "kept" if t == s11 else "via withdrawn " + s11))
wb2 = openpyxl.load_workbook(os.path.join(N, "sp800-53r5-to-iso-27001-mapping-2022-OLIR-2023-10-12-UPDATED.xlsx"), read_only=True)
out2 = set()
for ws in wb2.worksheets:
    if ws.title == "Definitions": continue
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not isinstance(r[0], str) or not isinstance(r[3], str): continue
        for ref in re.split(r"[,;\n]\s*", r[3]):
            m = re.match(r"A\.?([5-8]\.\d+)", ref.strip())
            if m: out2.add((base(r[0]), m.group(1)))
with open(os.path.join(ROOT, "data", "library", "nist_csf2_80053.csv"), "w", newline="") as fh:
    w = csv.writer(fh); w.writerow(["csf2", "sp80053r5", "basis"]); w.writerows(sorted(out1))
with open(os.path.join(ROOT, "data", "library", "nist_80053_iso27002.csv"), "w", newline="") as fh:
    w = csv.writer(fh); w.writerow(["sp80053r5", "iso27002"]); w.writerows(sorted(out2))
print("CSF 1.1 subcategories:", dict(stats), "-> CSF 2.0 to 800-53 pairs:", len(out1), "on", len({a for a, b, c in out1}), "subcategories")
print("800-53 to ISO 27002 (Annex A) pairs:", len(out2), "on", len({a for a, b in out2}), "controls")
