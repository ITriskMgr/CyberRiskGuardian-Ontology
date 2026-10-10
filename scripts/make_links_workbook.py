# -*- coding: utf-8 -*-
"""0.3.0 review workbook: new term, framework links (approve by group, comment individually), case scenarios and
case-calibrated family estimates."""
import csv, json, os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "build", "CRGO-0.3.0-review.xlsx")
F = "Arial"; base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
hfill = PatternFill("solid", fgColor="1F3864"); hfont = Font(name=F, size=10, bold=True, color="FFFFFF")
inp = PatternFill("solid", fgColor="FFF2CC"); thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()

def sheet(title, cols, rows, inputs=("Decision", "Comment"), wrap_rows=True):
    ws = wb.create_sheet(title)
    cols = cols + [(c, 12 if c == "Decision" else 40) for c in inputs]
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill; c.alignment = wrap; c.border = box
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(list(r) + [None] * len(inputs), 1):
            if isinstance(v, float) and v != v: v = None
            c = ws.cell(row=i, column=j, value=v); c.font = base; c.border = box
            if wrap_rows: c.alignment = wrap
            if j > len(cols) - len(inputs): c.fill = inp
    last = len(rows) + 1
    if "Decision" in inputs:
        dv = DataValidation(type="list", formula1='"Approve,Reject,Revise"', allow_blank=True); ws.add_data_validation(dv)
        col = get_column_letter(len(cols) - len(inputs) + 1); dv.add("%s2:%s%d" % (col, col, last))
    ws.freeze_panes = "B2"; ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), last)
    return ws

terms = {t["key"]: t for t in json.load(open(os.path.join(ROOT, "build", "terms.json"), encoding="utf-8"))}
links = pd.read_csv(os.path.join(ROOT, "data", "library", "links.csv"), dtype=str)
fam = pd.read_csv(os.path.join(ROOT, "data", "library", "families.csv"))
names = dict(zip(fam.family, fam.name))
titles = {}
for line in open(os.path.join(ROOT, "data", "work", "linking", "reference.md"), encoding="utf-8"):
    if " | " in line:
        k, v = line.split(" | ")[:2]; titles[k.strip()] = v.strip()
SCHEMES = [("attack", "MITRE ATT&CK technique (on the threat event)"), ("csf2", "NIST CSF 2.0 subcategory"), ("sp80053r5", "NIST SP 800-53 control"),
           ("iso27002", "ISO/IEC 27002:2022 control"), ("loi25", "Loi 25 obligation")]
CONF = {"H": "high", "M": "medium", "L": "low"}

rm = wb.active; rm.title = "Read me"
L = [("CRGO 0.3.0 draft — scenario library with framework links and case estimates", Font(name=F, size=14, bold=True)), ("", base),
     ("Fill in only the yellow columns. Decision: Approve, Reject or Revise.", base), ("", base), ("Sheets", bold),
     ("New term — hasRelevantRequirement (scenario → framework requirement).", base),
     ("Links by group — approve the %d links in bulk by framework and confidence (15 groups); an independent check found high-confidence links 97%% correct and low ones 26%% correct." % len(links), base),
     ("Links — every link, filterable; put a comment on any single link to reject or change it (it overrides the group decision).", base),
     ("Case scenarios — the 67 case scenarios added to the library, with the families they refine.", base),
     ("Case estimates — the 273 families that received case-calibrated estimates.", base), ("", base),
     ("Suggested approach: approve high and medium groups, leave low groups proposed (or reject them), and spot-check a few families you know well.", base)]
for i, (t, f) in enumerate(L, 1):
    rm.cell(row=i, column=1, value=t).font = f
rm.column_dimensions["A"].width = 140

t = terms["hasRelevantRequirement"]
sheet("New term", [("Key", 22), ("Label (en)", 24), ("Label (fr)", 26), ("Definition (en)", 60), ("Definition (fr)", 60), ("Domain → range", 30)],
      [[t["key"], t["label_en"], t["label_fr"], t["def_en"], t["def_fr"], t["parent"]]])

grp = links.groupby(["scheme", "confidence"]).agg(n=("id", "size"), fams=("family", "nunique")).reset_index()
rows = []
for s, label in SCHEMES:
    for c in ("H", "M", "L"):
        g = grp[(grp.scheme == s) & (grp.confidence == c)]
        if len(g):
            rows.append([label, CONF[c], int(g.n.iloc[0]), int(g.fams.iloc[0])])
sheet("Links by group", [("Framework", 40), ("Confidence", 12), ("Links", 10), ("Families", 10)], rows)

lr = [[l.family, names.get(l.family, ""), dict(SCHEMES)[l.scheme].split(" (")[0], l.id, titles.get(l.id, ""), CONF[l.confidence], l.basis] for l in links.itertuples()]
sheet("Links", [("Family", 9), ("Scenario", 50), ("Framework", 26), ("ID", 22), ("Title", 55), ("Confidence", 11), ("Basis", 18)], lr,
      inputs=("Comment",), wrap_rows=False)

sc = json.load(open(os.path.join(ROOT, "data", "library", "case_scenarios.json"), encoding="utf-8"))
mt = {m["sid"]: m["families"] for m in map(json.loads, open(os.path.join(ROOT, "data", "library", "case_matches.jsonl"), encoding="utf-8"))}
P = ["PbA", "Pbx", "De", "Dm", "Th", "Mu", "red_p", "red_i", "cvss"]
sheet("Case scenarios", [("Case scenario", 22), ("Case", 30), ("Name", 45)] + [(p, 7) for p in P] + [("Refines families (grade)", 45)],
      [[x["sid"], x["case"], x["name"]] + [x["estimates"].get(p) for p in P] + ["; ".join("%s %s (%s)" % (m["family"], names.get(m["family"], "")[:40], m["m"]) for m in mt.get(x["sid"], []))] for x in sc])

per = {}
for x in sc:
    for m in mt.get(x["sid"], []):
        per.setdefault(m["family"], []).append((m["m"], x["sid"]))
sheet("Case estimates", [("Family", 9), ("Scenario", 50), ("Matched case scenarios (grade)", 60), ("Confidence", 11)],
      [[f, names.get(f, ""), "; ".join("%s (%s)" % (s_, g) for g, s_ in v), "medium" if any(g == "H" for g, _ in v) else "low"] for f, v in sorted(per.items())])
wb.save(out); print(out)
