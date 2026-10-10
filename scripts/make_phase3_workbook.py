# -*- coding: utf-8 -*-
"""Phase 3 review workbook: new and changed terms, scenario themes, category map, vocabularies and library families."""
import csv, json, os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "pipeline", "library"))
import themes as T  # noqa: E402
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "build", "CRGO-0.2.0-phase3-review.xlsx")

F = "Arial"
base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
hfill = PatternFill("solid", fgColor="1F3864"); hfont = Font(name=F, size=10, bold=True, color="FFFFFF")
inp = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()

def sheet(title, cols, rows, decision=True, wrap_rows=True):
    ws = wb.create_sheet(title)
    cols = cols + ([("Decision", 12), ("Comment", 40)] if decision else [("Comment", 40)])
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill; c.alignment = wrap; c.border = box
        ws.column_dimensions[get_column_letter(j)].width = w
    n_in = 2 if decision else 1
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(list(r) + [None] * n_in, 1):
            if isinstance(v, float) and v != v:
                v = None
            c = ws.cell(row=i, column=j, value=v); c.font = base; c.border = box
            if wrap_rows: c.alignment = wrap
            if j > len(cols) - n_in: c.fill = inp
    last = len(rows) + 1
    if decision:
        dv = DataValidation(type="list", formula1='"Approve,Reject,Revise"', allow_blank=True)
        ws.add_data_validation(dv); dc = get_column_letter(len(cols) - 1); dv.add("%s2:%s%d" % (dc, dc, last))
    ws.freeze_panes = "B2"; ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), last)
    return ws

terms = {t["key"]: t for t in json.load(open(os.path.join(ROOT, "build", "terms.json"), encoding="utf-8"))}
approved = {l.split("\t")[0] for l in open(os.path.join(ROOT, "src", "approvals.tsv"), encoding="utf-8") if l.strip() and not l.startswith("#")}
new = [k for k in terms if k not in approved and terms[k]["kind"] != "value set (class)"]
CONF = {"H": "high", "M": "medium", "L": "low"}
fam = pd.read_csv(os.path.join(ROOT, "data", "library", "families.csv"))
rows = pd.read_csv(os.path.join(ROOT, "data", "library", "rows.csv"))

# Read me
rm = wb.active; rm.title = "Read me"
L = [("CRGO 0.2.0 draft — Phase 3 review", Font(name=F, size=14, bold=True)), ("", base),
     ("Fill in only the yellow columns. Decision: Approve, Reject or Revise; Comment: what to change.", base), ("", base),
     ("Sheets", bold),
     ("New and changed terms — %d ontology terms added or reworded for Phase 3 (status proposed)." % len(new), base),
     ("Themes — the controlled list of %d scenario themes that replaces the 225 source categories." % len(T.THEMES), base),
     ("Category map — how each source category maps to a theme (used when the scenario name gives no theme).", base),
     ("Vocabularies — the framework vocabularies built for mapping, with source and licence.", base),
     ("Library families — the %d repaired scenario families (from %d usable rows); comment on any you want changed." % (len(fam), len(rows)), base),
     ("", base), ("Decisions taken on 2026-10-09: CIS left out; CCM controls private only; ISO titles published; COBIT 2019 kept (docs/decisions.md, section 7).", base),
     ("", base), ("Example of a filled-in row", bold), ("Themes · ThemeOther · Decision: Revise · Comment: split into 'governance gaps' and 'other'.", base)]
for i, (t, f) in enumerate(L, 1):
    rm.cell(row=i, column=1, value=t).font = f
rm.column_dimensions["A"].width = 140

sheet("New and changed terms", [("Module", 9), ("Kind", 14), ("Key", 22), ("IRI", 34), ("Label (en)", 24), ("Label (fr)", 24), ("Definition (en)", 55),
                                 ("Definition (fr)", 55), ("Parent / domain → range", 26), ("Confidence", 10), ("Rationale", 35)],
      [[terms[k][c] for c in ("module", "kind", "key", "iri", "label_en", "label_fr", "def_en", "def_fr", "parent")] +
       [CONF[terms[k]["confidence"]], terms[k]["rationale"] or ("Definition reworded (removed class EnvironmentalHazard)" if k == "hasThreatSourceType" else "")] for k in new])

th_rows = rows.groupby("theme").size(); th_fam = fam.groupby("theme").size()
sheet("Themes", [("Key", 18), ("Label (en)", 30), ("Label (fr)", 34), ("Definition (en)", 60), ("Definition (fr)", 60), ("Rows", 8), ("Families", 9)],
      [[k, en, fr, den, dfr, int(th_rows.get(k, 0)), int(th_fam.get(k, 0))] for k, en, fr, den, dfr, _ in T.THEMES])

cm = list(csv.DictReader(open(os.path.join(ROOT, "data", "library", "category_map.tsv"), encoding="utf-8"), delimiter="\t"))
sheet("Category map", [("Source category", 50), ("Rows", 8), ("Theme", 22)], [[r["source_category"], int(r["rows"]), r["theme"]] for r in cm],
      decision=False, wrap_rows=False)

VOC = [("theme", "CRGO scenario themes", 23, "CRGO", "CC BY-NC 4.0"),
       ("csf2", "NIST CSF 2.0 (functions, categories, subcategories; withdrawn kept as deprecated)", 225, "NIST OSCAL catalogue 1.2.0", "US Government work"),
       ("sp800-53r5", "NIST SP 800-53 Rev. 5.2.0 (families, controls, enhancements; titles only)", 1216, "NIST OSCAL catalogue", "US Government work"),
       ("iso27002-2022", "ISO/IEC 27002:2022 (93 IDs and English titles, 4 themes)", 97, "ISO table of contents, cross-checked with 4 public mappings", "(c) ISO/IEC; IDs and titles only, for mapping"),
       ("cobit2019", "COBIT 2019 (40 objectives, 5 domains)", 45, "ISACA COBIT 2019 introduction; SFIA mapping", "(c) ISACA; names only, for mapping"),
       ("loi25", "Québec Law 25 obligations (P-39.1 and C-1.1 sections, en and fr)", 30, "LégisQuébec, section by section", "CRGO paraphrases"),
       ("ccm41", "CSA CCM v4.1 (17 domains only)", 17, "AWS v4.1 guide, ISMS Core", "(c) CSA; controls not redistributable"),
       ("attack", "MITRE ATT&CK Enterprise 19.2 (tactics, techniques, sub-techniques)", 712, "MITRE attack-stix-data", "ATT&CK Terms of Use")]
sheet("Vocabularies", [("Scheme", 14), ("Content", 60), ("Concepts", 10), ("Source", 45), ("Licence", 40)], [list(v) for v in VOC])

cols = ["family", "name", "theme", "n_rows", "sources", "PbA", "Pbx", "De", "Dm", "Th", "Mu", "red_p", "red_i", "cvss", "cost",
        "crg_est", "crg_res", "crg_ratio", "crg_cls", "threat_source_types", "asset_categories", "impact_dimensions", "rows"]
W = {"name": 45, "rows": 30, "threat_source_types": 22, "asset_categories": 28, "impact_dimensions": 28}
fr = fam[cols].copy()
for c in ("crg_est", "crg_res"):
    fr[c] = fr[c].round(1)
fr["crg_ratio"] = fr["crg_ratio"].round(3)
sheet("Library families", [(c, W.get(c, 10)) for c in cols], fr.values.tolist(), decision=False, wrap_rows=False)
wb.save(out); print(out)
