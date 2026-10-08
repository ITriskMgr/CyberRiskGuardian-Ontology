# -*- coding: utf-8 -*-
"""Build the term review workbook from build/terms.json (run after src/build.py)."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "build", "CRGO-0.1.0-term-review.xlsx")
terms = json.load(open(os.path.join(ROOT, "build", "terms.json"), encoding="utf-8"))
MOD = {"gov": 0, "core": 1, "measure": 2, "crg": 3}
KIND = {"class": 0, "value set (class)": 1, "value": 2, "object property": 3, "data property": 4, "annotation property": 5, "individual": 6}
terms.sort(key=lambda t: (MOD[t["module"]], KIND.get(t["kind"], 9)))
CONF = {"H": "high", "M": "medium", "L": "low"}

F = "Arial"
base, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
hfill = PatternFill("solid", fgColor="1F3864"); hfont = Font(name=F, size=10, bold=True, color="FFFFFF")
inp = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
# --- Read me
rm = wb.active; rm.title = "Read me"
lines = [
    ("CyberRiskGuardian Ontology 0.1.0 — term review (Phase 2)", Font(name=F, size=14, bold=True)),
    ("", base),
    ("Every term below is in status 'proposed'. Nothing becomes 'approved' until you mark it here and the decision is applied.", base),
    ("", base),
    ("How to review", bold),
    ("1. Go to the Terms sheet. Only the yellow columns (Decision, Comment) are for you to fill in.", base),
    ("2. Decision: choose Approve, Reject or Revise from the drop-down. Leave blank if not reviewed yet.", base),
    ("3. Comment: required for Revise and Reject — say what should change (label, definition, parent, French wording, etc.).", base),
    ("4. Filter by Module, Kind or Confidence to review in batches. Medium- and low-confidence terms are the ones most worth your attention.", base),
    ("5. Summary shows progress; it updates as you fill in decisions.", base),
    ("", base),
    ("Example of a filled-in row", bold),
    ("Key: ControlMeasure  ·  Decision: Revise  ·  Comment: Relabel EN to 'control implementation record'; keep FR.", base),
    ("", base),
    ("Columns", bold),
    ("Module — gov, core, measure or crg.  Kind — class, value set, value (member of a value set), property, individual.", base),
    ("Key — the readable name used in the source specification; the IRI is opaque and stable.", base),
    ("Parent / domain → range — superclass for classes, set for values, domain and range for properties ('—' = left open on purpose).", base),
    ("Confidence — high / medium / low, with the rationale for anything not obviously settled.", base),
    ("Sources — the evidence the term is grounded in (standards are paraphrased; no text reproduced).", base),
    ("", base),
    ("Also see docs/decisions.md in the repository for the design decisions, the review fixes and the deferred items.", base),
]
for i, (t, f) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=t); c.font = f
rm.column_dimensions["A"].width = 130

# --- Terms
ws = wb.create_sheet("Terms")
cols = [("Module", 9), ("Kind", 14), ("Key", 24), ("IRI", 36), ("Label (en)", 28), ("Label (fr)", 28),
        ("Definition (en)", 60), ("Definition (fr)", 60), ("Parent / domain → range", 30), ("Confidence", 11),
        ("Confidence rationale", 40), ("Sources", 40), ("Notes", 30), ("Decision", 12), ("Comment", 45)]
for j, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill; c.alignment = wrap; c.border = box
    ws.column_dimensions[get_column_letter(j)].width = w
for i, t in enumerate(terms, 2):
    row = [t["module"], t["kind"], t["key"], t["iri"], t["label_en"], t["label_fr"], t["def_en"], t["def_fr"],
           t["parent"], CONF.get(t["confidence"], t["confidence"]), t["rationale"], t["sources"], t["extra"], None, None]
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = base; c.alignment = wrap; c.border = box
        if j >= 14: c.fill = inp
n = len(terms) + 1
dv = DataValidation(type="list", formula1='"Approve,Reject,Revise"', allow_blank=True,
                    error="Choose Approve, Reject or Revise", errorTitle="Decision")
ws.add_data_validation(dv); dv.add("N2:N%d" % n)
ws.conditional_formatting.add("J2:J%d" % n, FormulaRule(formula=['$J2="low"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add("J2:J%d" % n, FormulaRule(formula=['$J2="medium"'], fill=PatternFill("solid", fgColor="FCE4D6")))
ws.freeze_panes = "D2"; ws.auto_filter.ref = "A1:O%d" % n

# --- Summary (formulas)
sm = wb.create_sheet("Summary")
hdr = ["Module", "Terms", "Approve", "Reject", "Revise", "Not reviewed", "Medium / low confidence"]
for j, h in enumerate(hdr, 1):
    c = sm.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill; c.border = box
R = "Terms!$A$2:$A$%d" % n; D = "Terms!$N$2:$N$%d" % n; C = "Terms!$J$2:$J$%d" % n
for i, m in enumerate(["gov", "core", "measure", "crg"], 2):
    sm.cell(row=i, column=1, value=m)
    sm.cell(row=i, column=2, value='=COUNTIF(%s,A%d)' % (R, i))
    for j, d in zip((3, 4, 5), ("Approve", "Reject", "Revise")):
        sm.cell(row=i, column=j, value='=COUNTIFS(%s,$A%d,%s,"%s")' % (R, i, D, d))
    sm.cell(row=i, column=6, value='=B%d-C%d-D%d-E%d' % (i, i, i, i))
    sm.cell(row=i, column=7, value='=COUNTIFS(%s,$A%d,%s,"medium")+COUNTIFS(%s,$A%d,%s,"low")' % (R, i, C, R, i, C))
sm.cell(row=6, column=1, value="Total")
for j in range(2, 8):
    L = get_column_letter(j); sm.cell(row=6, column=j, value="=SUM(%s2:%s5)" % (L, L))
for r in sm.iter_rows(min_row=2, max_row=6, max_col=7):
    for c in r: c.font = bold if c.row == 6 else base; c.border = box
sm.cell(row=8, column=1, value="Counts update from the Decision column of the Terms sheet.").font = Font(name=F, size=9, italic=True)
sm.column_dimensions["A"].width = 12
for j in range(2, 8): sm.column_dimensions[get_column_letter(j)].width = 16
sm.column_dimensions["G"].width = 24
wb.save(out); print(out, len(terms))
