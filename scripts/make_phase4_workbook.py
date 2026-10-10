# -*- coding: utf-8 -*-
"""Phase 4 review workbook: competency questions, commonly assessed risks, synonyms, proposed changes."""
import glob, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "pipeline", "query"))
import identify as I
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "build", "CRGO-0.4.0-phase4-review.xlsx")
F = "Arial"; base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
hfill = PatternFill("solid", fgColor="1F3864"); hfont = Font(name=F, size=10, bold=True, color="FFFFFF")
inp = PatternFill("solid", fgColor="FFF2CC"); thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()
def sheet(title, cols, rows):
    ws = wb.create_sheet(title); cols = cols + [("Decision", 12), ("Comment", 45)]
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill; c.alignment = wrap; c.border = box
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(list(r) + [None, None], 1):
            c = ws.cell(row=i, column=j, value=v); c.font = base; c.alignment = wrap; c.border = box
            if j > len(cols) - 2: c.fill = inp
    dv = DataValidation(type="list", formula1='"Approve,Reject,Revise"', allow_blank=True); ws.add_data_validation(dv)
    col = get_column_letter(len(cols) - 1); dv.add("%s2:%s%d" % (col, col, len(rows) + 1)); ws.freeze_panes = "B2"
rm = wb.active; rm.title = "Read me"
L = [("CRGO 0.4.0 draft — Phase 4 review", Font(name=F, size=14, bold=True)), ("", base),
     ("Fill in only the yellow columns. No ontology term changed in Phase 4; these are the tests, the identification vocabulary and proposed changes.", base), ("", base),
     ("Competency questions — the 20 questions the ontology must answer (all pass). Approve them as the test suite, or ask for others.", base),
     ("Commonly assessed risks — the 18 risks the identification sweep always looks for (my choice; adjust freely).", base),
     ("Synonyms — keyword expansion used by identification (my choice; adjust freely).", base),
     ("Proposed changes — library additions and ontology refinements suggested by the evaluation and the pitfall scan.", base), ("", base),
     ("Blind test on MediBec: recall 0.52 in the first truly blind run; 0.62 after fixes, but not independent (see docs/decisions.md 10.2).", bold)]
for i, (t, f) in enumerate(L, 1): rm.cell(row=i, column=1, value=t).font = f
rm.column_dimensions["A"].width = 140
cq = []
for f in sorted(glob.glob(os.path.join(ROOT, "tests", "cq", "CQ*.rq"))):
    t = open(f, encoding="utf-8").read()
    cq.append([os.path.basename(f)[:-3], re.search(r"\((\w+)\)", t).group(1), re.search(r"# Question: (.*)", t).group(1), re.search(r"# Expect: (.*)", t).group(1), "pass"])
sheet("Competency questions", [("ID", 7), ("Kind", 14), ("Question", 80), ("Expected result", 30), ("Result", 8)], cq)
sheet("Commonly assessed risks", [("Risk", 36), ("Words matched in scenario names", 70)], [[n, ", ".join(w)] for n, w in I.BASELINE])
sheet("Synonyms", [("Group", 90)], [[", ".join(g)] for g in I.SYNONYMS])
prop = [
 ["Library", "Add three generic scenario families the library lacks: cloud tenant takeover through token theft or malicious app consent; staff pasting confidential data into generative-AI tools; interception or tampering of HL7/DICOM clinical messaging.", "Found missing in the blind test; written by Claude, proposed, no estimates until calibrated."],
 ["Library", "Promote case scenarios that match no family well (e.g. MediBec S20, S26, S29) into generic families automatically.", "Keeps the generic library growing from assessed cases."],
 ["Ontology", "Declare the value-set classes pairwise disjoint (a confidence level is never a threat source type, etc.).", "Pitfall P10; safe, improves error detection."],
 ["Ontology", "Declare LibraryScenario and AssessmentScenario disjoint.", "Pitfall P10; an assessment scenario refines a library scenario rather than being one."],
 ["Identification", "Measure recall on a case that played no part in tuning (for example a new teaching case, or an ACRB or CanaRetail assessment once one is done in CyberRiskGuardian).", "The MediBec figure after tuning is not independent."],
 ["Identification", "Add the Québec health-information legislation (for private medical centres) to the obligations vocabulary.", "Raised in the first blind test; only P-39.1 is mapped today."],
 ["Data", "Fill missing Loi 25 links for scenarios that reach personal information without being tagged with the privacy impact dimension.", "Raised by the link review and the blind test."],
]
sheet("Proposed changes", [("Area", 14), ("Proposal", 90), ("Why", 60)], prop)
wb.save(out); print(out)
